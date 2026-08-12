from pathlib import Path

from openpyxl import load_workbook

from app.service.extractor.document_text_extractor import ExtractedTextPage


class ExcelTextExtractor:
    # Excel 시트의 셀 값을 행 단위 텍스트로 변환합니다.

    def supports(self, extension: str, mime_type: str | None = None) -> bool:
        return extension.lower().lstrip(".") in {"xlsx", "xlsm"}

    def extract(self, file_path: Path) -> list[ExtractedTextPage]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            pages: list[ExtractedTextPage] = []

            for index, sheet in enumerate(workbook.worksheets, start=1):
                rows: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        rows.append(" | ".join(values))

                text = "\n".join(rows).strip()
                if text:
                    pages.append(ExtractedTextPage(index, sheet.title, text))

            return pages
        finally:
            # read_only 워크북은 내부 zip 핸들을 명시적으로 닫아야 한다 —
            # 안 닫으면 Windows에서 임시 파일 삭제(unlink)가 PermissionError로 실패한다.
            workbook.close()